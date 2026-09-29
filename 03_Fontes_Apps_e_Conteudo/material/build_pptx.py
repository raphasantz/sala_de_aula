# -*- coding: utf-8 -*-
"""
Gerador da APRESENTACAO PowerPoint — Programação para Internet I (Turma 2/2026).
Slides para uso em sala de aula: teoria por aula, exercícios e testes rápidos
(com gabarito nas notas do apresentador).
Uso: python3 build_pptx.py
"""
import os
import re
import sys

_HERE = os.path.dirname(os.path.abspath(__file__))
_ROOT = os.path.dirname(_HERE)
sys.path.insert(0, _HERE)

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

import importlib as _il
_PKG = _il.import_module(os.environ.get("CONTEUDO_PKG", "conteudo"))
MODULOS = _PKG.MODULOS
MODULOS_POR_NUM = _PKG.MODULOS_POR_NUM
comuns = _PKG.comuns
avaliacoes = _PKG.avaliacoes

# ---------------------------------------------------------------------------
# PALETA
# ---------------------------------------------------------------------------
NAVY    = RGBColor(0x0D, 0x2B, 0x4E)
NAVY2   = RGBColor(0x14, 0x3A, 0x66)
PRIMARY = RGBColor(0x1B, 0x5F, 0xAA)
SKY     = RGBColor(0xE9, 0xF1, 0xFA)
SKY_TXT = RGBColor(0xBB, 0xD4, 0xEE)
AMBER   = RGBColor(0xF5, 0x9E, 0x0B)
AMBER_D = RGBColor(0x9A, 0x62, 0x06)
AMBER_BG= RGBColor(0xFE, 0xF4, 0xE2)
INK     = RGBColor(0x1F, 0x29, 0x37)
SOFT    = RGBColor(0x51, 0x60, 0x6F)
LINE    = RGBColor(0xC9, 0xD8, 0xE8)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
CODE_BG = RGBColor(0x0F, 0x22, 0x37)
CODE_FG = RGBColor(0xDC, 0xE9, 0xF7)
GREEN   = RGBColor(0x1E, 0x7B, 0x34)
RED     = RGBColor(0xC0, 0x39, 0x2B)

# fontes (padrão Windows/PowerPoint)
FH = "Segoe UI"          # headings
FB = "Segoe UI"          # body
FC = "Consolas"          # código

# cores de sintaxe no slide de código
SC_COM = RGBColor(0x7F, 0x96, 0xAD)
SC_STR = RGBColor(0x9E, 0xCE, 0x6A)
SC_KEY = RGBColor(0xD3, 0x8A, 0xE0)
SC_VAR = RGBColor(0xF2, 0x77, 0x7A)
SC_FUN = RGBColor(0x6C, 0xB6, 0xFF)
SC_NUM = RGBColor(0xE5, 0xC0, 0x7B)

SW = Inches(13.333)
SH = Inches(7.5)

ROMAN = {1: "I", 2: "II", 3: "III", 4: "IV", 5: "V", 6: "VI"}

# ---------------------------------------------------------------------------
# realce de sintaxe -> segmentos (texto, cor)
# ---------------------------------------------------------------------------
PHP_WORDS = ("echo|print|if|else|elseif|endif|for|endforeach|foreach|while|do|switch|case|"
             "break|default|function|return|true|false|null|exit|die|include|require|"
             "include_once|require_once|define|as|new|isset|empty")
RE_PHP_TOK = re.compile(r"(\$[A-Za-z_]\w*)|\b(" + PHP_WORDS + r")\b|\b([A-Za-z_]\w*)(?=\s*\()")
RE_STR = re.compile(r'("[^"]*"|\'[^\']*\')')
RE_COM = re.compile(r"(//.*$|#.*$)")
RE_TAG = re.compile(r"(<\/?[A-Za-z][^>]*>)")
RE_PTAG = re.compile(r"(<\?php\b|<\?=|\?>)")

def _segs_php(line):
    segs = []
    m = RE_COM.search(line)
    code, com = (line[:m.start()], line[m.start():]) if m else (line, None)
    pos = 0
    for mp in RE_PTAG.finditer(code):
        segs += _segs_php_core(code[pos:mp.start()])
        segs.append((mp.group(1), SC_KEY)); pos = mp.end()
    segs += _segs_php_core(code[pos:])
    if com:
        segs.append((com, SC_COM))
    return segs

def _segs_php_core(txt):
    segs = []
    pos = 0
    for ms in RE_STR.finditer(txt):
        segs += _segs_php_tok(txt[pos:ms.start()])
        segs.append((ms.group(1), SC_STR)); pos = ms.end()
    segs += _segs_php_tok(txt[pos:])
    return segs

def _segs_php_tok(txt):
    segs = []
    pos = 0
    for m in RE_PHP_TOK.finditer(txt):
        if m.start() > pos:
            segs.append((txt[pos:m.start()], CODE_FG))
        if m.group(1):
            segs.append((m.group(1), SC_VAR))
        elif m.group(2):
            segs.append((m.group(2), SC_KEY))
        else:
            segs.append((m.group(3), SC_FUN))
        pos = m.end()
    if pos < len(txt):
        segs.append((txt[pos:], CODE_FG))
    return segs

def _segs_html(line):
    segs = []
    pos = 0
    for m in RE_TAG.finditer(line):
        if m.start() > pos:
            segs.append((line[pos:m.start()], CODE_FG))
        tag = m.group(1)
        # strings dentro da tag em verde
        tpos = 0
        for ms in RE_STR.finditer(tag):
            if ms.start() > tpos:
                segs.append((tag[tpos:ms.start()], SC_FUN))
            segs.append((ms.group(1), SC_STR)); tpos = ms.end()
        if tpos < len(tag):
            segs.append((tag[tpos:], SC_FUN))
        pos = m.end()
    if pos < len(line):
        segs.append((line[pos:], CODE_FG))
    return segs

def segs_for(line, lang):
    if lang == "html":
        return _segs_html(line)
    if lang == "php":
        return _segs_php(line)
    if line.strip().startswith("#"):
        return [(line, SC_COM)]
    return [(line, CODE_FG)]

# ---------------------------------------------------------------------------
# HELPERS DE SLIDE
# ---------------------------------------------------------------------------
def add_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])

def rect(slide, x, y, w, h, color, rounded=False, line_color=None, line_w=None, radius=0.06):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        x, y, w, h)
    if rounded:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line_color is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line_color
        shp.line.width = line_w or Pt(1)
    shp.shadow.inherit = False
    return shp

def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tb, tf

def setp(p, text, size=18, color=INK, bold=False, italic=False, font=FB,
         align=PP_ALIGN.LEFT, space_after=6, space_before=0, line=None):
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    if line:
        p.line_spacing = line
    r = p.add_run()
    r.text = text
    f = r.font
    f.size = Pt(size)
    f.color.rgb = color
    f.bold = bold
    f.italic = italic
    f.name = font
    return r

def para_runs(p, segs, size=14, font=FC, align=PP_ALIGN.LEFT, space_after=2, line=None):
    p.alignment = align
    p.space_after = Pt(space_after)
    if line:
        p.line_spacing = line
    for txt, col in segs:
        r = p.add_run()
        r.text = txt
        r.font.size = Pt(size)
        r.font.color.rgb = col
        r.font.name = font
        r.font.bold = False

def clean_txt(txt):
    """Remove marcação HTML de textos de conteúdo para uso em slide."""
    t = txt
    t = re.sub(r"<b>(.*?)</b>", r"\1", t)
    t = re.sub(r"<i>(.*?)</i>", r"\1", t)
    t = (t.replace("&lt;", "<").replace("&gt;", ">").replace("&amp;", "&")
          .replace("&quot;", '"').replace("<br>", " "))
    t = re.sub(r"<[^>]+>", "", t)
    return t

def footer(slide, prs, idx, label=""):
    rect(slide, Inches(0), SH - Inches(0.42), SW, Inches(0.42), SKY)
    rect(slide, Inches(0), SH - Inches(0.46), SW, Inches(0.045), LINE)
    tb, tf = textbox(slide, Inches(0.55), SH - Inches(0.38), Inches(9.5), Inches(0.34))
    setp(tf.paragraphs[0], comuns.META["pptx_footer"]
         + (("  ·  " + label) if label else ""), size=9, color=SOFT, space_after=0)
    rect(slide, SW - Inches(1.12), SH - Inches(0.40), Inches(0.36), Inches(0.36), NAVY,
         rounded=True, radius=0.5)
    tb2, tf2 = textbox(slide, SW - Inches(1.12), SH - Inches(0.40), Inches(0.36), Inches(0.36),
                       anchor=MSO_ANCHOR.MIDDLE)
    setp(tf2.paragraphs[0], str(idx), size=11, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER, space_after=0)

def header(slide, kicker, title, accent=AMBER):
    rect(slide, Inches(0), Inches(0), SW, Inches(0.09), NAVY)
    rect(slide, Inches(0), Inches(0.09), Inches(4.3), Inches(0.05), accent)
    tb, tf = textbox(slide, Inches(0.62), Inches(0.34), Inches(12.1), Inches(0.4))
    setp(tf.paragraphs[0], kicker.upper(), size=11.5, color=AMBER_D, bold=True, space_after=0)
    tb2, tf2 = textbox(slide, Inches(0.62), Inches(0.68), Inches(12.1), Inches(0.9))
    setp(tf2.paragraphs[0], title, size=27, color=NAVY, bold=True, space_after=0)

def notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = clean_txt(text)

# ---------------------------------------------------------------------------
# TIPOS DE SLIDE
# ---------------------------------------------------------------------------
def slide_titulo(prs, idx):
    s = add_slide(prs)
    rect(s, 0, 0, SW, SH, NAVY)
    rect(s, 0, 0, SW, Inches(0.22), AMBER)
    # círculos
    rect(s, SW - Inches(2.6), Inches(-1.4), Inches(4.4), Inches(4.4), NAVY2, rounded=True, radius=0.5)
    rect(s, SW - Inches(1.5), Inches(-0.5), Inches(2.6), Inches(2.6), RGBColor(0x1B, 0x4B, 0x85),
         rounded=True, radius=0.5)
    rect(s, Inches(-1.2), SH - Inches(2.2), Inches(3.6), Inches(3.6), NAVY2, rounded=True, radius=0.5)
    tb, tf = textbox(s, Inches(0.9), Inches(0.85), Inches(11.5), Inches(0.5))
    setp(tf.paragraphs[0], comuns.META["kicker"], size=13, color=AMBER, bold=True, space_after=0)
    tb, tf = textbox(s, Inches(0.9), Inches(1.5), Inches(11.6), Inches(2.2))
    setp(tf.paragraphs[0], comuns.META["pptx_capa1"], size=54, color=WHITE, bold=True, space_after=0)
    p = tf.add_paragraph(); setp(p, comuns.META["pptx_capa2"], size=54, color=WHITE, bold=True, space_after=0)
    rect(s, Inches(0.92), Inches(3.62), Inches(3.1), Inches(0.075), AMBER)
    tb, tf = textbox(s, Inches(0.9), Inches(3.95), Inches(11.4), Inches(0.9))
    setp(tf.paragraphs[0], comuns.META["pptx_sub"], size=17, color=SKY_TXT, space_after=0)
    # chips de informação
    chips = [("72 aulas", "24 semanas · 3/semana"),
             ("60 h + 10 h", "presenciais + AVA"),
             ("12 módulos", "6 partes · 3 avaliações")]
    x = Inches(0.9)
    for a, b in chips:
        shp = rect(s, x, Inches(5.05), Inches(3.6), Inches(0.95), NAVY2, rounded=True, radius=0.12)
        tb, tf = textbox(s, x + Inches(0.25), Inches(5.2), Inches(3.1), Inches(0.7))
        setp(tf.paragraphs[0], a, size=16, color=WHITE, bold=True, space_after=0)
        p = tf.add_paragraph(); setp(p, b, size=11.5, color=SKY_TXT, space_after=0)
        x += Inches(3.85)
    tb, tf = textbox(s, Inches(0.9), Inches(6.55), Inches(11.5), Inches(0.5))
    setp(tf.paragraphs[0], comuns.META["pptx_rodape"], size=10.5, color=RGBColor(0x9F, 0xBB, 0xD9), space_after=0)
    return s

def slide_front(prs, idx, kicker, title, bullets, sub=None, two_col=False):
    s = add_slide(prs)
    rect(s, 0, 0, SW, SH, WHITE)
    header(s, kicker, title)
    y = Inches(1.75)
    if sub:
        tb, tf = textbox(s, Inches(0.62), y, Inches(12.1), Inches(0.5))
        setp(tf.paragraphs[0], sub, size=14, color=SOFT, italic=True, space_after=0)
        y += Inches(0.55)
    tb, tf = textbox(s, Inches(0.62), y, Inches(12.1), SH - y - Inches(0.7))
    first = True
    for b in bullets:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        setp(p, "•  " + clean_txt(b), size=16.5, color=INK, space_after=9, line=1.05)
    footer(s, prs, idx)
    return s

def slide_part(prs, idx, parte, mods):
    s = add_slide(prs)
    rect(s, 0, 0, SW, SH, NAVY)
    rect(s, 0, SH - Inches(0.16), SW, Inches(0.16), AMBER)
    rect(s, SW - Inches(2.4), Inches(-1.2), Inches(4.2), Inches(4.2), NAVY2, rounded=True, radius=0.5)
    tb, tf = textbox(s, Inches(0.9), Inches(0.95), Inches(11), Inches(1.3))
    setp(tf.paragraphs[0], f"PARTE {ROMAN[parte['num']]}", size=52, color=WHITE, bold=True, space_after=0)
    rect(s, Inches(0.92), Inches(2.15), Inches(2.4), Inches(0.07), AMBER)
    tb, tf = textbox(s, Inches(0.9), Inches(2.5), Inches(11.5), Inches(1.1))
    setp(tf.paragraphs[0], parte["titulo"].upper(), size=26, color=WHITE, bold=True, space_after=0)
    chip = rect(s, Inches(0.92), Inches(3.75), Inches(6.4), Inches(0.62), AMBER, rounded=True, radius=0.3)
    tb, tf = textbox(s, Inches(1.2), Inches(3.87), Inches(6.0), Inches(0.4))
    setp(tf.paragraphs[0], f"{parte['aulas'].upper()}  ·  {parte['semanas'].upper()}",
         size=13, color=NAVY, bold=True, space_after=0)
    y = Inches(4.8)
    tb, tf = textbox(s, Inches(0.92), y, Inches(11.4), Inches(2.2))
    first = True
    for n, t in mods:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        setp(p, f"Módulo {n} — {t}", size=15, color=SKY_TXT, space_after=6)
    tb, tf = textbox(s, SW - Inches(2.2), SH - Inches(0.75), Inches(1.5), Inches(0.4))
    setp(tf.paragraphs[0], str(idx), size=12, color=WHITE, bold=True, align=PP_ALIGN.RIGHT, space_after=0)
    return s

def slide_module(prs, idx, m):
    s = add_slide(prs)
    rect(s, 0, 0, SW, SH, NAVY)
    rect(s, Inches(0), Inches(0), Inches(0.28), SH, AMBER)
    rect(s, SW - Inches(2.6), SH - Inches(2.6), Inches(4.6), Inches(4.6), NAVY2, rounded=True, radius=0.5)
    tb, tf = textbox(s, Inches(0.95), Inches(1.05), Inches(11.5), Inches(0.5))
    setp(tf.paragraphs[0], f"MÓDULO {m['num']}  ·  {m['aulas_faixa'].upper()}  ·  {m['semanas'].upper()}",
         size=14, color=AMBER, bold=True, space_after=0)
    tb, tf = textbox(s, Inches(0.95), Inches(1.7), Inches(11.4), Inches(1.9))
    setp(tf.paragraphs[0], m["titulo"], size=36, color=WHITE, bold=True, space_after=0)
    tb, tf = textbox(s, Inches(0.95), Inches(3.6), Inches(11.3), Inches(0.5))
    setp(tf.paragraphs[0], f"PARTE {ROMAN[m['parte_num']]} — {m['parte_titulo']}",
         size=14, color=SKY_TXT, space_after=0)
    tb, tf = textbox(s, Inches(0.95), Inches(4.5), Inches(11.3), Inches(2.3))
    first = True
    for o in m["objetivos"]:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        setp(p, "•  " + clean_txt(o), size=14.5, color=SKY_TXT, space_after=7)
    tb, tf = textbox(s, SW - Inches(2.2), SH - Inches(0.75), Inches(1.5), Inches(0.4))
    setp(tf.paragraphs[0], str(idx), size=12, color=WHITE, bold=True, align=PP_ALIGN.RIGHT, space_after=0)
    return s

def slide_teoria(prs, idx, m, aula, parte_label):
    s = add_slide(prs)
    rect(s, 0, 0, SW, SH, WHITE)
    header(s, f"MÓDULO {m['num']} · AULA {aula['num']}", aula["titulo"])
    pontos = aula["slides"]["pontos"]
    tb, tf = textbox(s, Inches(0.62), Inches(1.8), Inches(12.1), Inches(4.9))
    first = True
    for b in pontos:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        setp(p, "•  " + clean_txt(b), size=18, color=INK, space_after=12, line=1.05)
    footer(s, prs, idx, parte_label)
    notes(s, aula["slides"].get("nota"))
    return s

def slide_codigo(prs, idx, m, aula, spec, parte_label):
    s = add_slide(prs)
    rect(s, 0, 0, SW, SH, WHITE)
    header(s, f"MÓDULO {m['num']} · AULA {aula['num']} · CÓDIGO", spec.get("titulo", "Exemplo"))
    linhas = spec["linhas"]
    n = len(linhas)
    top = Inches(1.75)
    h = min(Inches(5.0), Inches(0.30 * n + 0.6))
    panel = rect(s, Inches(0.62), top, Inches(12.1), h, CODE_BG, rounded=True, radius=0.045)
    tb, tf = textbox(s, Inches(0.95), top + Inches(0.25), Inches(11.5), h - Inches(0.4))
    first = True
    for ln in linhas:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        segs = segs_for(ln, spec.get("ling", "php")) if ln.strip() else [(" ", CODE_FG)]
        para_runs(p, segs, size=13.5, space_after=1, line=1.02)
    footer(s, prs, idx, parte_label)
    notes(s, aula["slides"].get("nota"))
    return s

def slide_extras(prs, idx, m, aula, parte_label):
    s = add_slide(prs)
    rect(s, 0, 0, SW, SH, WHITE)
    header(s, f"MÓDULO {m['num']} · AULA {aula['num']}", aula["titulo"] + " (cont.)")
    tb, tf = textbox(s, Inches(0.62), Inches(1.8), Inches(12.1), Inches(4.9))
    first = True
    for b in aula["slides"]["extra"]:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        setp(p, "•  " + clean_txt(b), size=17, color=INK, space_after=10)
    footer(s, prs, idx, parte_label)
    notes(s, aula["slides"].get("nota"))
    return s

def slide_exercicios(prs, idx, m, parte_label):
    s = add_slide(prs)
    rect(s, 0, 0, SW, SH, WHITE)
    rect(s, Inches(0), Inches(0), SW, Inches(0.09), NAVY)
    rect(s, Inches(0), Inches(0.09), Inches(4.3), Inches(0.05), AMBER)
    # faixa
    faixa = rect(s, Inches(0.62), Inches(0.42), Inches(12.1), Inches(0.95), AMBER_BG, rounded=True, radius=0.12)
    rect(s, Inches(0.62), Inches(0.42), Inches(0.14), Inches(0.95), AMBER)
    tb, tf = textbox(s, Inches(1.0), Inches(0.55), Inches(11.5), Inches(0.75))
    setp(tf.paragraphs[0], f"EXERCÍCIOS PRÁTICOS — MÓDULO {m['num']}", size=19, color=AMBER_D,
         bold=True, space_after=2)
    p = tf.add_paragraph()
    setp(p, "Resolva no laboratório · complete em casa pelo AVA", size=12, color=SOFT, space_after=0)
    y = Inches(1.75)
    tb, tf = textbox(s, Inches(0.62), y, Inches(12.1), Inches(5.0))
    first = True
    for ex in m["exercicios"]:
        tipo = {"pratico": "PRÁTICO", "escrito": "ESCRITO", "grupo": "GRUPO"}[ex["tipo"]]
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        setp(p, f"{ex['num']}.  {clean_txt(ex['titulo'])}", size=16.5, color=NAVY, bold=True,
             space_after=2)
        p2 = tf.add_paragraph()
        setp(p2, f"     [{tipo}]  " + clean_txt(ex["enunciado"])[:210], size=13, color=SOFT,
             space_after=10, line=1.03)
    footer(s, prs, idx, parte_label)
    return s

def slide_quiz(prs, idx, m, q, qnum, parte_label):
    s = add_slide(prs)
    rect(s, 0, 0, SW, SH, WHITE)
    rect(s, Inches(0), Inches(0), SW, Inches(0.09), NAVY)
    rect(s, Inches(0), Inches(0.09), Inches(4.3), Inches(0.05), AMBER)
    faixa = rect(s, Inches(0.62), Inches(0.40), Inches(12.1), Inches(0.72), PRIMARY, rounded=True, radius=0.16)
    tb, tf = textbox(s, Inches(0.95), Inches(0.53), Inches(11.5), Inches(0.5))
    setp(tf.paragraphs[0], f"TESTE RÁPIDO — MÓDULO {m['num']}   ·   QUESTÃO {qnum}/5",
         size=15, color=WHITE, bold=True, space_after=0)
    tb, tf = textbox(s, Inches(0.62), Inches(1.42), Inches(12.1), Inches(1.3))
    setp(tf.paragraphs[0], clean_txt(q["enunciado"]), size=20, color=INK, bold=True, space_after=0,
         line=1.08)
    letters = ["A", "B", "C", "D"]
    y = Inches(2.85)
    for li, alt in enumerate(q["alt"]):
        chip = rect(s, Inches(0.62), y, Inches(0.52), Inches(0.52), SKY, rounded=True, radius=0.25)
        tb, tf = textbox(s, Inches(0.62), y + Inches(0.08), Inches(0.52), Inches(0.4))
        setp(tf.paragraphs[0], letters[li], size=16, color=PRIMARY, bold=True,
             align=PP_ALIGN.CENTER, space_after=0)
        tb, tf = textbox(s, Inches(1.35), y + Inches(0.06), Inches(11.3), Inches(0.75))
        setp(tf.paragraphs[0], clean_txt(alt), size=16.5, color=INK, space_after=0, line=1.02)
        y += Inches(0.95)
    tb, tf = textbox(s, Inches(0.62), SH - Inches(0.92), Inches(12.0), Inches(0.4))
    setp(tf.paragraphs[0], "Resposta do professor: ver notas do apresentador (F5 → anotações).",
         size=10.5, color=SOFT, italic=True, space_after=0)
    footer(s, prs, idx, parte_label)
    notes(s, f"RESPOSTA: {letters[q['resposta']]} — {q.get('comentario','')}")
    return s

def slide_aviso_avaliacao(prs, idx, m, ref, parte_label):
    s = add_slide(prs)
    rect(s, 0, 0, SW, SH, WHITE)
    header(s, f"MÓDULO {m['num']}", "Avaliação à vista")
    box = rect(s, Inches(0.62), Inches(1.8), Inches(12.1), Inches(3.6), AMBER_BG, rounded=True,
               radius=0.06)
    rect(s, Inches(0.62), Inches(1.8), Inches(0.16), Inches(3.6), AMBER)
    tb, tf = textbox(s, Inches(1.1), Inches(2.1), Inches(11.3), Inches(3.1))
    txts = {
        "A1": [("AVALIAÇÃO 1 — 30 pts (semanas 1 a 8)", True),
               ("Trabalho prático de HTML “site da empresa fictícia” — 10 pts (entrega semana 6)", False),
               ("Atividades extraclasse/AVA e exercícios — 5 pts", False),
               ("Teste escrito-prático sobre Web e HTML — 15 pts (semana 8)", False),
               ("Modelo do teste e rubrica do trabalho: seção AVALIAÇÕES da apostila.", False)],
        "A1-teste": [("SEMANA 8 — TESTE ESCRITO-PRÁTICO DA A1 (15 pts)", True),
                     ("Conteúdo: Web (cliente/servidor, URL, DNS, hospedagem) + HTML (módulos 3 a 5)", False),
                     ("Parte A: 6 objetivas · Parte B: 3 práticas (escrever/corrigir código)", False),
                     ("Modelo completo na seção AVALIAÇÕES da apostila — use como revisão!", False)],
        "A2": [("AVALIAÇÃO 2 — 30 pts (semana 16 · aula 48)", True),
               ("Prova obrigatória — pode ser em duplas ou com consulta (Plano de Ensino)", False),
               ("HTML (estrutura a formulários) · PHP (variáveis a arrays/funções) · GET/POST/validação", False),
               ("~60% prática (leitura/escrita de código) + ~40% conceitual", False),
               ("Modelo de prova na seção AVALIAÇÕES da apostila.", False)],
        "A3": [("AVALIAÇÃO 3 — 40 pts (semana 24)", True),
               ("Prova escrita final individual e cumulativa — 30 pts (aulas 70–71)", False),
               ("Projeto integrador: entregas parciais + apresentação — 10 pts (aula 72)", False),
               ("Rubrica do projeto na seção AVALIAÇÕES da apostila.", False)],
    }
    first = True
    for t, b in txts[ref]:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        setp(p, ("•  " if not b else "") + t, size=17 if b else 15, color=AMBER_D if b else INK,
             bold=b, space_after=9)
    footer(s, prs, idx, parte_label)
    return s

def slide_projeto(prs, idx):
    PF = avaliacoes.PROJETO_FINAL
    s = add_slide(prs)
    rect(s, 0, 0, SW, SH, WHITE)
    header(s, "AULAS 49 A 72 · SEMANAS 17 A 24", "Projeto Final Integrador")
    box = rect(s, Inches(0.62), Inches(1.7), Inches(6.0), Inches(4.9), SKY, rounded=True, radius=0.05)
    tb, tf = textbox(s, Inches(0.95), Inches(1.95), Inches(5.4), Inches(4.5))
    setp(tf.paragraphs[0], "O QUE É", size=13, color=PRIMARY, bold=True, space_after=4)
    setp(tf.add_paragraph(), clean_txt(PF["o_que_e"]), size=13.5, color=INK, space_after=10, line=1.1)
    setp(tf.add_paragraph(), "TEMAS SUGERIDOS", size=13, color=PRIMARY, bold=True, space_after=4)
    for t in PF["temas"]:
        setp(tf.add_paragraph(), "•  " + t, size=13, color=INK, space_after=3)
    box = rect(s, Inches(6.85), Inches(1.7), Inches(5.85), Inches(4.9), WHITE, rounded=True,
               radius=0.05, line_color=LINE, line_w=Pt(1))
    tb, tf = textbox(s, Inches(7.15), Inches(1.95), Inches(5.3), Inches(4.5))
    setp(tf.paragraphs[0], "ENTREGAS (6)", size=13, color=PRIMARY, bold=True, space_after=4)
    for e in PF["entregas"]:
        setp(tf.add_paragraph(), f"{e[0]} — {clean_txt(e[1])[:70]}", size=12, color=INK, space_after=5)
    setp(tf.add_paragraph(), "CRITÉRIOS (10 pts)", size=13, color=PRIMARY, bold=True, space_after=4)
    for a, b, c in PF["criterios_10pts"]:
        setp(tf.add_paragraph(), f"•  {a} ({b})", size=12, color=INK, space_after=3)
    footer(s, prs, idx, "PARTE VI")
    return s

def slide_encerramento(prs, idx):
    s = add_slide(prs)
    rect(s, 0, 0, SW, SH, NAVY)
    rect(s, 0, 0, SW, Inches(0.22), AMBER)
    tb, tf = textbox(s, Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.6))
    setp(tf.paragraphs[0], "Bons estudos e bom código!", size=44, color=WHITE, bold=True, space_after=0)
    rect(s, Inches(0.92), Inches(3.6), Inches(2.6), Inches(0.07), AMBER)
    tb, tf = textbox(s, Inches(0.9), Inches(4.0), Inches(11.4), Inches(1.4))
    setp(tf.paragraphs[0], comuns.META["pptx_encerra1"], size=16, color=SKY_TXT, space_after=6)
    setp(tf.add_paragraph(), comuns.META["pptx_encerra2"], size=14, color=SKY_TXT, space_after=0)
    tb, tf = textbox(s, SW - Inches(2.2), SH - Inches(0.75), Inches(1.5), Inches(0.4))
    setp(tf.paragraphs[0], str(idx), size=12, color=WHITE, bold=True, align=PP_ALIGN.RIGHT, space_after=0)
    return s

# ---------------------------------------------------------------------------
# MONTAGEM
# ---------------------------------------------------------------------------
def build(out_path):
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    idx = 0

    def nxt():
        nonlocal idx
        idx += 1
        return idx

    slide_titulo(prs, nxt())

    f0 = comuns.FRONT[0]
    slide_front(prs, nxt(), f0["kick"], f0["tit"], f0["bullets"], sub=f0["sub"])

    f1 = comuns.FRONT[1]
    slide_front(prs, nxt(), f1["kick"], f1["tit"], f1["bullets"])

    f2 = comuns.FRONT[2]
    slide_front(prs, nxt(), f2["kick"], f2["tit"], f2["bullets"])

    f3 = comuns.FRONT[3]
    slide_front(prs, nxt(), f3["kick"], f3["tit"], f3["bullets"], sub=f3.get("sub"))

    ordem_partes = comuns.ORDEM_PARTES
    for parte in comuns.PARTES:
        mods = ordem_partes[parte["num"]]
        slide_part(prs, nxt(), parte, [(n, MODULOS_POR_NUM[n]["titulo"]) for n in mods])
        for n in mods:
            m = MODULOS_POR_NUM[n]
            parte_label = f"PARTE {ROMAN[m['parte_num']]}"
            slide_module(prs, nxt(), m)
            for aula in m["aulas"]:
                sl = aula.get("slides", {})
                slide_teoria(prs, nxt(), m, aula, parte_label)
                if sl.get("codigo"):
                    slide_codigo(prs, nxt(), m, aula, sl["codigo"], parte_label)
                if sl.get("extra"):
                    slide_extras(prs, nxt(), m, aula, parte_label)
            slide_exercicios(prs, nxt(), m, parte_label)
            for qi, q in enumerate(m["teste_rapido"], 1):
                slide_quiz(prs, nxt(), m, q, qi, parte_label)
            if m.get("avaliacao_ref"):
                slide_aviso_avaliacao(prs, nxt(), m, m["avaliacao_ref"], parte_label)

    slide_projeto(prs, nxt())
    slide_encerramento(prs, nxt())

    prs.save(out_path)
    print("OK:", out_path, "·", idx, "slides")

if __name__ == "__main__":
    build(os.path.join(_ROOT, "output", f"Apresentacao_{comuns.META['slug']}.pptx"))

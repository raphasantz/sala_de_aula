# -*- coding: utf-8 -*-
"""Preview aproximado de PPTX (para QA de layout) via PIL."""
import os, sys
from PIL import Image, ImageDraw, ImageFont

_HERE = os.path.dirname(os.path.abspath(__file__))
_ROOT = os.path.dirname(_HERE)
sys.path.insert(0, _HERE)

from pptx import Presentation
from pptx.util import Emu

FD = os.path.join(_ROOT, "fonts")
_cache = {}
def font(size, bold=False, italic=False, mono=False):
    key = (round(size), bold, italic, mono)
    if key not in _cache:
        if mono:
            path = os.path.join(FD, "JetBrainsMono-Bold.ttf" if bold else "JetBrainsMono-Regular.ttf")
        elif bold:
            path = os.path.join(FD, "Inter-Bold.ttf")
        elif italic:
            path = os.path.join(FD, "Inter-Italic.ttf")
        else:
            path = os.path.join(FD, "Inter-Regular.ttf")
        _cache[key] = ImageFont.truetype(path, max(6, int(round(size))))
    return _cache[key]

SCALE = 100  # px por polegada
W, H = int(13.333 * SCALE), int(7.5 * SCALE)

def rgb(c):
    try:
        return (c[0], c[1], c[2])
    except Exception:
        return (0, 0, 0)

def shape_fill_color(shp):
    try:
        f = shp.fill
        if f.type is not None and str(f.type) == "MSO_FILL.SOLID (1)" or f.type == 1:
            return rgb(f.fore_color.rgb)
    except Exception:
        pass
    return None

def wrap(draw, text, fnt, maxw):
    words = text.split()
    lines = []
    cur = ""
    for w in words:
        t = (cur + " " + w).strip()
        if draw.textlength(t, font=fnt) <= maxw or not cur:
            cur = t
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines

def render_slide(slide, path):
    img = Image.new("RGB", (W, H), (255, 255, 255))
    d = ImageDraw.Draw(img)
    def emu(v):
        return int(Emu(v).inches * SCALE) if v is not None else 0
    shapes = list(slide.shapes)
    for shp in shapes:
        x, y = emu(shp.left), emu(shp.top)
        w, h = emu(shp.width), emu(shp.height)
        # forma com preenchimento
        col = None
        try:
            if shp.fill.type == 1:
                col = rgb(shp.fill.fore_color.rgb)
        except Exception:
            col = None
        if col is not None and shp.shape_type is not None and shp.has_text_frame is False:
            d.rectangle([x, y, x + w, y + h], fill=col)
        elif col is not None and not shp.text_frame.text.strip():
            d.rectangle([x, y, x + w, y + h], fill=col)
        if shp.has_text_frame:
            tf = shp.text_frame
            yy = y + 2
            for p in tf.paragraphs:
                runs = p.runs
                if not runs:
                    yy += 8
                    continue
                r0 = runs[0]
                size = (r0.font.size.pt if r0.font.size else 18) * SCALE / 72.0
                bold = bool(r0.font.bold)
                colr = rgb(r0.font.color.rgb) if r0.font.color and r0.font.color.type is not None else (31, 41, 55)
                mono = (r0.font.name or "") in ("Consolas", "Mono")
                fnt = font(size, bold=bold, mono=mono)
                text = "".join(r.text for r in runs)
                lines = wrap(d, text, fnt, max(20, w - 4))
                lh = size * 1.28
                for ln in lines:
                    if yy > H - 2:
                        break
                    # alinhamento
                    tx = x
                    try:
                        from pptx.enum.text import PP_ALIGN
                        if p.alignment == PP_ALIGN.CENTER:
                            tx = x + (w - d.textlength(ln, font=fnt)) / 2
                        elif p.alignment == PP_ALIGN.RIGHT:
                            tx = x + w - d.textlength(ln, font=fnt)
                    except Exception:
                        pass
                    d.text((tx, yy), ln, font=fnt, fill=colr)
                    yy += lh
                yy += (p.space_after.pt if p.space_after else 0) * SCALE / 72.0
    img.save(path)

if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("pages", nargs="+", type=int)
    ap.add_argument("--out", default="preview")
    a = ap.parse_args()
    prs = Presentation(a.pptx)
    os.makedirs(a.out, exist_ok=True)
    slides = list(prs.slides)
    for p in a.pages:
        render_slide(slides[p - 1], os.path.join(a.out, f"pptx_{p:03d}.png"))
    print("renderizado:", a.pages)
